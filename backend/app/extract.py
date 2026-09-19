"""Turn uploaded documents (PDF, Word, PowerPoint, Excel) into readable text.

Only text-bearing content is extracted: paragraphs, headings, tables, slide text,
speaker notes and cell values. Images inside documents are ignored, and a PDF
with no text layer (a scan) is rejected rather than stored empty.
"""
import io
import re
from typing import Any

SUPPORTED = {"pdf", "docx", "pptx", "xlsx"}
MAX_TEXT = 300_000
MAX_PDF_PAGES = 500
MAX_SHEET_ROWS = 300
MAX_SHEET_COLS = 40


class ExtractError(Exception):
    """A message that is safe and useful to show to the person uploading."""


def extension(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def _tidy(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\x00", "")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _md_row(cells: list[str]) -> str:
    return "| " + " | ".join(c.replace("|", "\\|").replace("\n", " ").strip() for c in cells) + " |"


def _md_table(rows: list[list[str]]) -> str:
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    lines = [_md_row(rows[0]), "| " + " | ".join("---" for _ in range(width)) + " |"]
    lines += [_md_row(r) for r in rows[1:]]
    return "\n".join(lines)


def _pdf(data: bytes) -> tuple[str, dict[str, Any]]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ExtractError("This PDF is password protected. Remove the password and upload it again.")
        pages = len(reader.pages)
        parts: list[str] = []
        for n, page in enumerate(reader.pages[:MAX_PDF_PAGES], 1):
            text = _tidy(page.extract_text() or "")
            if text:
                parts.append(f"[Page {n}]\n{text}")
    except ExtractError:
        raise
    except (PdfReadError, ValueError, KeyError, OSError) as exc:
        raise ExtractError("That PDF could not be read. It may be damaged.") from exc
    body = "\n\n".join(parts)
    if len(body) < 20:
        raise ExtractError(
            "This PDF has no selectable text (it looks like a scan). Reading scanned pages needs text recognition, which is not supported yet."
        )
    note = f"\n\n(Only the first {MAX_PDF_PAGES} of {pages} pages were read.)" if pages > MAX_PDF_PAGES else ""
    return body + note, {"pages": pages}


def _docx(data: bytes) -> tuple[str, dict[str, Any]]:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001 - python-docx raises several unrelated types on bad files
        raise ExtractError("That Word file could not be read. Only .docx files are supported, not the older .doc format.") from exc
    out: list[str] = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, doc)
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            if style == "Title":
                out.append(f"# {text}")
            elif style.startswith("Heading"):
                m = re.search(r"(\d)", style)
                level = min(int(m.group(1)), 4) if m else 2
                out.append(f"{'#' * level} {text}")
            elif "List" in style:
                out.append(f"- {text}")
            else:
                out.append(text)
        elif tag == "tbl":
            table = Table(child, doc)
            md = _md_table([[c.text for c in row.cells] for row in table.rows])
            if md:
                out.append(md)
    return _tidy("\n\n".join(out)), {}


def _pptx(data: bytes) -> tuple[str, dict[str, Any]]:
    from pptx import Presentation

    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ExtractError("That PowerPoint file could not be read. Only .pptx files are supported, not the older .ppt format.") from exc
    out: list[str] = []
    for n, slide in enumerate(prs.slides, 1):
        title = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text.strip()
        out.append(f"## Slide {n}" + (f": {title}" if title else ""))
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if shape.has_text_frame:
                lines = [p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()]
                if lines:
                    out.append("\n".join(f"- {l}" for l in lines))
            elif getattr(shape, "has_table", False) and shape.has_table:
                md = _md_table([[c.text for c in row.cells] for row in shape.table.rows])
                if md:
                    out.append(md)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                out.append(f"Speaker notes: {notes}")
    return _tidy("\n\n".join(out)), {"pages": len(prs.slides)}


def _xlsx(data: bytes) -> tuple[str, dict[str, Any]]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ExtractError("That Excel file could not be read. Only .xlsx files are supported, not the older .xls format.") from exc
    out: list[str] = []
    for ws in wb.worksheets:
        rows: list[list[str]] = []
        total = 0
        for row in ws.iter_rows(values_only=True):
            total += 1
            if total > MAX_SHEET_ROWS + 1:
                continue
            rows.append(["" if v is None else str(v) for v in row[:MAX_SHEET_COLS]])
        md = _md_table(rows)
        if not md:
            continue
        note = f"\n\n(Showing the first {MAX_SHEET_ROWS} of {total - 1} rows.)" if total - 1 > MAX_SHEET_ROWS else ""
        out.append(f"## Sheet: {ws.title}\n\n{md}{note}")
    wb.close()
    return _tidy("\n\n".join(out)), {"pages": len(wb.sheetnames)}


def extract_text(filename: str, data: bytes) -> tuple[str, dict[str, Any]]:
    ext = extension(filename)
    if ext not in SUPPORTED:
        raise ExtractError("Supported documents are PDF, Word (.docx), PowerPoint (.pptx) and Excel (.xlsx).")
    text, meta = {"pdf": _pdf, "docx": _docx, "pptx": _pptx, "xlsx": _xlsx}[ext](data)
    if not text.strip():
        raise ExtractError("No readable text was found in that file.")
    if len(text) > MAX_TEXT:
        text = text[:MAX_TEXT].rstrip() + "\n\n(The rest of the document was cut off to keep it manageable.)"
        meta["truncated"] = True
    return text, meta
